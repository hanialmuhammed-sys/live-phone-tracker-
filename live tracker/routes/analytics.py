import csv
import io
from datetime import datetime, timedelta, time

from flask import Blueprint, render_template, jsonify, Response, send_file
from flask_login import login_required, current_user

from models import Location, Visit
from utils import haversine_distance

analytics_bp = Blueprint("analytics", __name__, url_prefix="/analytics")

MOVING_THRESHOLD_MPS = 0.5  # ~1.8 km/h — below this, we call it "stationary"


def _day_bounds(day):
    return datetime.combine(day, time.min), datetime.combine(day, time.max)


def compute_daily_stats(user_id, day=None):
    """Part 3 — distance, speed, and moving/stationary time for one day."""
    day = day or datetime.utcnow().date()
    start, end = _day_bounds(day)

    locations = (
        Location.query
        .filter(Location.user_id == user_id)
        .filter(Location.timestamp >= start, Location.timestamp <= end)
        .order_by(Location.timestamp)
        .all()
    )

    total_distance_m = 0.0
    moving_seconds = 0.0
    stationary_seconds = 0.0
    max_speed_mps = 0.0
    moving_speed_samples = []

    for prev, curr in zip(locations, locations[1:]):
        distance_m = haversine_distance(
            prev.latitude, prev.longitude,
            curr.latitude, curr.longitude
        )
        time_delta = (curr.timestamp - prev.timestamp).total_seconds()

        if time_delta <= 0:
            continue

        # Prefer the device's own reported speed; fall back to distance/time
        segment_speed = curr.speed if curr.speed is not None else (distance_m / time_delta)

        total_distance_m += distance_m
        max_speed_mps = max(max_speed_mps, segment_speed)

        if segment_speed >= MOVING_THRESHOLD_MPS:
            moving_seconds += time_delta
            moving_speed_samples.append(segment_speed)
        else:
            stationary_seconds += time_delta

    average_speed_mps = (
        sum(moving_speed_samples) / len(moving_speed_samples)
        if moving_speed_samples else 0.0
    )

    return {
        "distance": round(total_distance_m / 1000, 1),      # km
        "average_speed": round(average_speed_mps * 3.6),     # km/h
        "max_speed": round(max_speed_mps * 3.6),              # km/h
        "moving_time": int(moving_seconds),                    # seconds
        "stationary_time": int(stationary_seconds),             # seconds
    }


def compute_weekly_distance(user_id):
    """Part 5 — distance per day for the current week (Mon-Sun)."""
    today = datetime.utcnow().date()
    week_start = today - timedelta(days=today.weekday())

    days = []
    for i in range(7):
        day = week_start + timedelta(days=i)
        stats = compute_daily_stats(user_id, day)
        days.append({"day": day.strftime("%a"), "distance": stats["distance"]})

    max_distance = max((d["distance"] for d in days), default=0) or 1
    for d in days:
        d["height_pct"] = round((d["distance"] / max_distance) * 100)

    return days


def compute_speed_series(user_id, day=None):
    """Part 5 — speed samples through the day, for a simple line chart."""
    day = day or datetime.utcnow().date()
    start, end = _day_bounds(day)

    locations = (
        Location.query
        .filter(Location.user_id == user_id)
        .filter(Location.timestamp >= start, Location.timestamp <= end)
        .order_by(Location.timestamp)
        .all()
    )

    series = []
    for prev, curr in zip(locations, locations[1:]):
        time_delta = (curr.timestamp - prev.timestamp).total_seconds()
        if time_delta <= 0:
            continue

        if curr.speed is not None:
            speed_mps = curr.speed
        else:
            distance_m = haversine_distance(
                prev.latitude, prev.longitude,
                curr.latitude, curr.longitude
            )
            speed_mps = distance_m / time_delta

        series.append({
            "time": curr.timestamp.strftime("%H:%M"),
            "speed": round(speed_mps * 3.6, 1)  # km/h
        })

    return series


def get_visits_for_day(user_id, day=None):
    """Part 4 — places stayed at for 10+ minutes."""
    day = day or datetime.utcnow().date()
    start, end = _day_bounds(day)

    visits = (
        Visit.query
        .filter(Visit.user_id == user_id)
        .filter(Visit.arrival_time >= start, Visit.arrival_time <= end)
        .order_by(Visit.arrival_time)
        .all()
    )

    result = []
    for visit in visits:
        result.append({
            "place": visit.place_name,
            "arrival": visit.arrival_time.strftime("%H:%M"),
            "departure": visit.departure_time.strftime("%H:%M"),
            "duration_seconds": visit.duration_seconds
        })

    return result


@analytics_bp.route("")
@login_required
def analytics_page():
    return render_template(
        "analytics.html",
        stats=compute_daily_stats(current_user.id),
        week=compute_weekly_distance(current_user.id),
        visits=get_visits_for_day(current_user.id)
    )


@analytics_bp.route("/today")
@login_required
def today():
    return jsonify(compute_daily_stats(current_user.id))


@analytics_bp.route("/week")
@login_required
def week():
    return jsonify(compute_weekly_distance(current_user.id))


@analytics_bp.route("/speeds")
@login_required
def speeds():
    return jsonify(compute_speed_series(current_user.id))


@analytics_bp.route("/places")
@login_required
def places():
    return jsonify(get_visits_for_day(current_user.id))


@analytics_bp.route("/export/csv")
@login_required
def export_csv():
    locations = (
        Location.query
        .filter_by(user_id=current_user.id)
        .order_by(Location.timestamp)
        .all()
    )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["timestamp", "latitude", "longitude", "speed_mps", "heading", "altitude", "battery"])

    for loc in locations:
        writer.writerow([
            loc.timestamp.isoformat(),
            loc.latitude,
            loc.longitude,
            loc.speed,
            loc.heading,
            loc.altitude,
            loc.battery
        ])

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=travel_history.csv"}
    )


@analytics_bp.route("/export/xlsx")
@login_required
def export_xlsx():
    from openpyxl import Workbook

    locations = (
        Location.query
        .filter_by(user_id=current_user.id)
        .order_by(Location.timestamp)
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Travel History"
    ws.append(["Timestamp", "Latitude", "Longitude", "Speed (m/s)", "Heading", "Altitude", "Battery"])

    for loc in locations:
        ws.append([
            loc.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            loc.latitude,
            loc.longitude,
            loc.speed,
            loc.heading,
            loc.altitude,
            loc.battery
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="travel_history.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@analytics_bp.route("/export/pdf")
@login_required
def export_pdf():
    from fpdf import FPDF

    stats = compute_daily_stats(current_user.id)
    locations = (
        Location.query
        .filter_by(user_id=current_user.id)
        .order_by(Location.timestamp)
        .all()
    )

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Travel History Report", ln=True)

    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"User: {current_user.username}", ln=True)
    pdf.cell(0, 8, f"Distance today: {stats['distance']} km", ln=True)
    pdf.cell(0, 8, f"Average speed: {stats['average_speed']} km/h", ln=True)
    pdf.cell(0, 8, f"Max speed: {stats['max_speed']} km/h", ln=True)
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(50, 8, "Timestamp", border=1)
    pdf.cell(35, 8, "Latitude", border=1)
    pdf.cell(35, 8, "Longitude", border=1)
    pdf.cell(30, 8, "Speed (m/s)", border=1)
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for loc in locations[-200:]:  # cap rows so long histories don't blow up the PDF
        pdf.cell(50, 7, loc.timestamp.strftime("%Y-%m-%d %H:%M:%S"), border=1)
        pdf.cell(35, 7, f"{loc.latitude:.5f}", border=1)
        pdf.cell(35, 7, f"{loc.longitude:.5f}", border=1)
        pdf.cell(30, 7, f"{loc.speed:.1f}" if loc.speed is not None else "-", border=1)
        pdf.ln()

    pdf_bytes = pdf.output()
    if isinstance(pdf_bytes, str):
        pdf_bytes = pdf_bytes.encode("latin-1")

    buffer = io.BytesIO(bytes(pdf_bytes))

    return send_file(
        buffer,
        as_attachment=True,
        download_name="travel_history.pdf",
        mimetype="application/pdf"
    )